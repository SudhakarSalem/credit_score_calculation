"""Excel (openpyxl) used as the data store for customers and credit score details."""
import threading
from datetime import datetime

from django.conf import settings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_lock = threading.Lock()
SHEET = "Customers"

COLUMNS = [
    "customer_id", "name", "email", "phone", "dob", "city",
    "on_time_payments", "total_payments", "late_30", "late_60", "late_90", "collections", "bankruptcy",
    "total_balance", "total_credit_limit",
    "oldest_account_years", "newest_account_months", "avg_account_age_years",
    "revolving_accounts", "installment_accounts", "hard_inquiries_12m", "new_accounts_12m",
    "local_score", "local_tier", "gemini_score", "gemini_tier", "gemini_risk", "gemini_breakdown",
    "calculated_at", "created_at",
]

NUMERIC = {
    "on_time_payments", "total_payments", "late_30", "late_60", "late_90", "collections",
    "total_balance", "total_credit_limit", "oldest_account_years", "newest_account_months",
    "avg_account_age_years", "revolving_accounts", "installment_accounts",
    "hard_inquiries_12m", "new_accounts_12m", "local_score", "gemini_score",
}


def _path():
    p = settings.EXCEL_PATH
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


def _style(ws):
    head_fill = PatternFill("solid", start_color="1F3864")
    for i, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(name) + 4)
    ws.column_dimensions[get_column_letter(COLUMNS.index("gemini_breakdown") + 1)].width = 40
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"


def _open():
    p = _path()
    if p.exists():
        wb = load_workbook(p)
        return wb, wb[SHEET]
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    _style(ws)
    return wb, ws


def _save(wb, ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
    wb.save(_path())


def _convert(col, v):
    if v is None or v == "":
        return None if col in NUMERIC else ""
    if col in NUMERIC:
        try:
            f = float(v)
            return int(f) if f.is_integer() else f
        except (TypeError, ValueError):
            return None
    return str(v)


def all_customers():
    with _lock:
        wb, ws = _open()
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            rows.append({col: _convert(col, v) for col, v in zip(COLUMNS, r)})
        return rows


def get_customer(cid):
    for c in all_customers():
        if c["customer_id"] == cid:
            return c
    return None


def next_id(rows):
    nums = [int(r["customer_id"][1:]) for r in rows if str(r["customer_id"]).startswith("C")]
    return f"C{(max(nums) if nums else 1000) + 1}"


def add_customer(data: dict) -> dict:
    with _lock:
        wb, ws = _open()
        existing = [{"customer_id": r[0]} for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        data = dict(data)
        data["customer_id"] = next_id(existing)
        data["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        ws.append([data.get(col) for col in COLUMNS])
        _save(wb, ws)
        return data


def update_customer(cid, fields: dict):
    with _lock:
        wb, ws = _open()
        for row in ws.iter_rows(min_row=2):
            if row[0].value == cid:
                for col, val in fields.items():
                    row[COLUMNS.index(col)].value = val
                _save(wb, ws)
                return True
    return False


def bulk_write(rows):
    """Replace all rows (used by the sample data generator and batch scoring)."""
    with _lock:
        p = _path()
        if p.exists():
            p.unlink()
        wb, ws = _open()
        for d in rows:
            ws.append([d.get(col) for col in COLUMNS])
        _save(wb, ws)
